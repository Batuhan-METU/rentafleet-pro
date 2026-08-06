"""Excel (.xlsx) raporları için ortak çalışma sayfası oluşturucu.

Tüm dışa aktarma servisleri (filo, kiralama, bakım) aynı görsel dili
kullanır: koyu dolgulu dondurulmuş başlık satırı, filtre okları, sayı/para
formatları ve içeriğe göre auto-fit sütun genişlikleri.
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CURRENCY_FORMAT = '#,##0.00" ₺"'
KM_FORMAT = "#,##0"
PERCENT_FORMAT = '0.0"%"'
DATE_FORMAT = "DD.MM.YYYY"

_MIN_COLUMN_WIDTH = 11.0
_MAX_COLUMN_WIDTH = 42.0

_HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _display_length(value: object, number_format: str | None) -> int:
    """Hücrenin Excel'de kaplayacağı yaklaşık karakter genişliği."""
    if value is None:
        return 0
    if isinstance(value, (datetime, date)):
        return len(DATE_FORMAT)
    if isinstance(value, (int, float)) and number_format == CURRENCY_FORMAT:
        return len(f"{value:,.2f} ₺")
    if isinstance(value, (int, float)) and number_format == KM_FORMAT:
        return len(f"{value:,.0f}")
    return len(str(value))


def build_worksheet_bytes(
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    *,
    number_formats: dict[int, str] | None = None,
) -> bytes:
    """Tek sayfalık .xlsx içeriği (bytes) üretir.

    ``number_formats``: 1 tabanlı sütun numarası → Excel sayı formatı.
    Tarih/datetime değerleri otomatik olarak ``DATE_FORMAT`` ile yazılır.
    """
    formats = number_formats or {}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    column_widths = [len(header) + 4 for header in headers]

    for row_index, row_values in enumerate(rows, start=2):
        for column, value in enumerate(row_values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            number_format = formats.get(column)
            if isinstance(value, (datetime, date)):
                cell.number_format = DATE_FORMAT
            elif number_format and value is not None:
                cell.number_format = number_format
            if isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal="right")

            width = _display_length(value, number_format) + 2
            if width > column_widths[column - 1]:
                column_widths[column - 1] = width

    # Sütunlar içeriğe göre genişletilir (auto-fit); okunabilirlik için
    # alt/üst sınırlar uygulanır.
    for column, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = min(
            max(float(width), _MIN_COLUMN_WIDTH), _MAX_COLUMN_WIDTH
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
