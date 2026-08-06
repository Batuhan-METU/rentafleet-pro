"""Gerçek filo araçları için operasyon gideri (bakım/lastik/hasar/muayene/
yıkama vb.) import eder (MALIYET_HESAPLAMA_METODOLOJISI §3.5'in girdisi).

Kaynak: ``datas/Arac_Bazli_Operasyon_Gideri.xlsx`` (sayfa ``Tüm Giderler``)
— plaka ile eşleştirilir, sadece bizim gerçek araçlarımıza karşılık gelen
satırlar aktarılır. Kaynağın "FYS Alt Kalem" sütunu ``ExpenseCategory``
seçeneklerine BİREBİR eşlenir (kategori listesi artık FYS'nin gerçek alt
kalem adlarıyla aynı).

Kullanım::

    python manage.py import_vehicle_operating_expenses ../datas/Arac_Bazli_Operasyon_Gideri.xlsx
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from typing import Any

from django.core.management.base import BaseCommand, CommandError, CommandParser
from django.db import transaction

from backend.apps.vehicle_finance.choices import ExpenseCategory
from backend.apps.vehicle_finance.models import VehicleExpense
from backend.apps.vehicles.models import Vehicle

SHEET_NAME = "Tüm Giderler"
COL_PLATE = "Plaka"
COL_DATE = "Tarih"
COL_SUB_ITEM = "FYS Alt Kalem"
COL_AMOUNT = "Net Tutar"

CATEGORY_MAP = {
    "Araç Mekanik": ExpenseCategory.VEHICLE_MECHANICAL,
    "Filo Tipi Değişikliği": ExpenseCategory.FLEET_TYPE_CHANGE,
    "Operasyon Birimi Sevk": ExpenseCategory.OPERATIONS_UNIT_DISPATCH,
    "Araç Bakım": ExpenseCategory.VEHICLE_MAINTENANCE,
    "SSH Masraf Yansıtma": ExpenseCategory.SSH_COST_REFLECTION,
    "Yazlık/Kışlık Lastik Değişim": ExpenseCategory.SEASONAL_TIRE_CHANGE,
    "Kayıp/Bozuk Anahtar": ExpenseCategory.LOST_OR_BROKEN_KEY,
    "Şüpheli Araç Bildirimi": ExpenseCategory.SUSPICIOUS_VEHICLE_REPORT,
    "Satın Alma": ExpenseCategory.PURCHASE,
    "Şube Araç Artırma Azaltma": ExpenseCategory.BRANCH_FLEET_REBALANCING,
    "Araç Hasar Onarım": ExpenseCategory.VEHICLE_DAMAGE_REPAIR,
    "ZF Muayene": ExpenseCategory.ZF_INSPECTION,
    "Plaka Ruhsat": ExpenseCategory.PLATE_AND_REGISTRATION,
    "Diğer Lastik İşlemleri": ExpenseCategory.OTHER_TIRE_SERVICES,
    "Müşavirlik Araç Muayene": ExpenseCategory.CONSULTANCY_VEHICLE_INSPECTION,
    "Lastik Hasar Onarım": ExpenseCategory.TIRE_DAMAGE_REPAIR,
    "Satış Birimi Sevk": ExpenseCategory.SALES_UNIT_DISPATCH,
    "Kiralama Dönüşü Yıkama": ExpenseCategory.RETURN_WASHING,
    "Kiralama Dönüşü Yakıt Alma": ExpenseCategory.RETURN_REFUELING,
    "Standart Yıkama": ExpenseCategory.STANDARD_WASHING,
    "Personel Araç Atama": ExpenseCategory.STAFF_VEHICLE_ASSIGNMENT,
    "Lastik Montaj (4 Lastik)": ExpenseCategory.TIRE_INSTALLATION_SET_OF_4,
    "Toplu İşlem": ExpenseCategory.BULK_OPERATION,
    "Diğer (Eşlenmemiş)": ExpenseCategory.OTHER,
}


def _parse_turkish_date(raw: object) -> date | None:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    if not raw:
        return None
    text = str(raw).strip().split(" ")[0]
    try:
        return datetime.strptime(text, "%d.%m.%Y").date()
    except ValueError:
        return None


class Command(BaseCommand):
    help = "Gerçek araçlara operasyon gideri import eder (plaka eşleştirmesiyle)."

    def add_arguments(self, parser: CommandParser) -> None:
        parser.add_argument("xlsx_path", type=str, help="Arac_Bazli_Operasyon_Gideri.xlsx yolu.")
        parser.add_argument(
            "--dry-run", action="store_true", help="Hiçbir şey kaydetmeden özet göster."
        )

    def handle(self, *args: Any, **options: Any) -> None:
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError("openpyxl kurulu değil.") from exc

        xlsx_path: str = options["xlsx_path"]
        try:
            workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        except (OSError, KeyError) as exc:
            raise CommandError(f"Dosya okunamadı: {xlsx_path} ({exc})") from exc

        if SHEET_NAME not in workbook.sheetnames:
            raise CommandError(f"'{SHEET_NAME}' sayfası bulunamadı.")
        sheet = workbook[SHEET_NAME]
        rows_iter = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
        for required in (COL_PLATE, COL_DATE, COL_SUB_ITEM, COL_AMOUNT):
            if required not in header:
                raise CommandError(f"Eksik sütun: {required}")
        idx = {h: i for i, h in enumerate(header)}

        our_vehicles = {v.plate_number: v for v in Vehicle.objects.all()}

        to_create: list[VehicleExpense] = []
        skipped_no_vehicle = 0
        skipped_bad_data = 0
        unmapped_categories: set[str] = set()

        for row in rows_iter:
            plate = str(row[idx[COL_PLATE]] or "").strip()
            vehicle = our_vehicles.get(plate)
            if vehicle is None:
                skipped_no_vehicle += 1
                continue

            sub_item = str(row[idx[COL_SUB_ITEM]] or "").strip()
            category = CATEGORY_MAP.get(sub_item, ExpenseCategory.OTHER)
            if sub_item not in CATEGORY_MAP:
                unmapped_categories.add(sub_item)

            amount = row[idx[COL_AMOUNT]]
            expense_date = _parse_turkish_date(row[idx[COL_DATE]])
            if amount is None or expense_date is None:
                skipped_bad_data += 1
                continue

            to_create.append(
                VehicleExpense(
                    vehicle=vehicle,
                    category=category,
                    custom_label=sub_item if category == ExpenseCategory.OTHER else "",
                    amount=Decimal(str(amount)),
                    expense_date=expense_date,
                )
            )

        self.stdout.write(f"Aktarıma hazır gider: {len(to_create)}")
        self.stdout.write(
            f"Atlanan: {skipped_no_vehicle} araç bizde yok, {skipped_bad_data} eksik veri."
        )
        if unmapped_categories:
            self.stdout.write(
                self.style.WARNING(f"Eşlenemeyen alt kalemler (OTHER'a düştü): {unmapped_categories}")
            )

        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("--dry-run: hiçbir şey kaydedilmedi."))
            return

        with transaction.atomic():
            VehicleExpense.objects.bulk_create(to_create)

        self.stdout.write(self.style.SUCCESS(f"{len(to_create)} gider aktarıldı."))
