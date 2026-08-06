"""Gerçek filo araçları için sigorta/kasko poliçeleri import eder
(MALIYET_HESAPLAMA_METODOLOJISI §3.4'ün girdisi).

Kaynak: ``datas/Sigorta_Kasko_Policeleri.xlsx`` (sayfa ``Detay``) — plaka
ile eşleştirilir, sadece bizim gerçek araçlarımıza karşılık gelen satırlar
aktarılır. Tutarlar zaten araç ölçeğinde (₺1.486-₺340.338 aralığı, filo-
geneli değil), ek ölçekleme gerekmez.

Kullanım::

    python manage.py import_vehicle_insurance_policies ../datas/Sigorta_Kasko_Policeleri.xlsx
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from typing import Any

from django.core.management.base import BaseCommand, CommandError, CommandParser
from django.db import transaction

from backend.apps.vehicle_finance.choices import InsurancePolicyType
from backend.apps.vehicle_finance.models import InsurancePolicy
from backend.apps.vehicles.models import Vehicle

SHEET_NAME = "Detay"
COL_PLATE = "Plaka"
COL_TYPE = "Poliçe Tipi"
COL_AMOUNT = "Poliçe Tutarı"
COL_START = "Başlangıç Tarihi"
COL_END = "Bitiş Tarihi"

POLICY_TYPE_MAP = {
    "Trafik Sigortası": InsurancePolicyType.LIABILITY,
    "Kasko Sigortası": InsurancePolicyType.COMPREHENSIVE,
    "Yeşil Kart Sigortası": InsurancePolicyType.GREEN_CARD,
}


def _parse_turkish_date(raw: object) -> date | None:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    if not raw:
        return None
    text = str(raw).strip().split(" ")[0]
    try:
        return datetime.strptime(text, "%d.%m.%Y").date()
    except ValueError:
        return None


class Command(BaseCommand):
    help = "Gerçek araçlara sigorta/kasko poliçesi import eder (plaka eşleştirmesiyle)."

    def add_arguments(self, parser: CommandParser) -> None:
        parser.add_argument("xlsx_path", type=str, help="Sigorta_Kasko_Policeleri.xlsx yolu.")
        parser.add_argument(
            "--dry-run", action="store_true", help="Hiçbir şey kaydetmeden özet göster."
        )

    def handle(self, *args: Any, **options: Any) -> None:
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError("openpyxl kurulu değil.") from exc

        xlsx_path: str = options["xlsx_path"]
        try:
            workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        except (OSError, KeyError) as exc:
            raise CommandError(f"Dosya okunamadı: {xlsx_path} ({exc})") from exc

        if SHEET_NAME not in workbook.sheetnames:
            raise CommandError(f"'{SHEET_NAME}' sayfası bulunamadı.")
        sheet = workbook[SHEET_NAME]
        rows_iter = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
        for required in (COL_PLATE, COL_TYPE, COL_AMOUNT, COL_START, COL_END):
            if required not in header:
                raise CommandError(f"Eksik sütun: {required}")
        idx = {h: i for i, h in enumerate(header)}

        our_vehicles = {v.plate_number: v for v in Vehicle.objects.all()}

        to_create: list[InsurancePolicy] = []
        skipped_no_vehicle = 0
        skipped_bad_type = 0
        skipped_bad_data = 0

        for row in rows_iter:
            plate = str(row[idx[COL_PLATE]] or "").strip()
            vehicle = our_vehicles.get(plate)
            if vehicle is None:
                skipped_no_vehicle += 1
                continue

            policy_type = POLICY_TYPE_MAP.get(str(row[idx[COL_TYPE]] or "").strip())
            if policy_type is None:
                skipped_bad_type += 1
                continue

            amount = row[idx[COL_AMOUNT]]
            start = _parse_turkish_date(row[idx[COL_START]])
            end = _parse_turkish_date(row[idx[COL_END]])
            if amount is None or start is None or end is None:
                skipped_bad_data += 1
                continue

            to_create.append(
                InsurancePolicy(
                    vehicle=vehicle,
                    policy_type=policy_type,
                    amount=Decimal(str(amount)),
                    start_date=start,
                    end_date=end,
                )
            )

        self.stdout.write(f"Aktarıma hazır poliçe: {len(to_create)}")
        self.stdout.write(
            f"Atlanan: {skipped_no_vehicle} araç bizde yok, "
            f"{skipped_bad_type} tanınmayan tip, {skipped_bad_data} eksik veri."
        )

        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("--dry-run: hiçbir şey kaydedilmedi."))
            return

        with transaction.atomic():
            InsurancePolicy.objects.bulk_create(to_create)

        self.stdout.write(self.style.SUCCESS(f"{len(to_create)} poliçe aktarıldı."))
