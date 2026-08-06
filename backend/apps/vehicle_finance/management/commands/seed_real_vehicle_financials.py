"""Gerçek filo araçları (`import_fleet_vehicles`) için finansal profil
üretir — MALIYET_HESAPLAMA_METODOLOJISI §2/§3.1/§3.2'nin zorunlu girdisi.
``VehicleFinancialProfile`` yokluğu, maliyet motorunun TÜM sonucu koşulsuz
``None`` döndürmesine sebep olan tek en büyük engeldir (bkz.
``total_cost_service._financing_missing_components``) — bu komut olmadan
gerçek araçların hiçbiri gerçek maliyet üretemez.

Kaynak: ``datas/Arac_Finansman_Tipi.xlsx`` (plaka → Finansman Tipi + Kredi/
Nakit tutar kırılımı). Bu dosya ZiraatFilo'nun kendi test ortamı için
SENTETİK üretilmiştir (bkz. metodoloji dokümanı §11) — biz bunu "gerçek
ZiraatFilo verisi" olarak değil, aynı vehicle evrenine (Vehicle_duzeltilmis.
xlsx) karşılık gelen, halihazırda tutarlı (kredi+nakit == satın alım
fiyatı) hazır bir finansman-tipi ataması olarak kullanıyoruz.

NOT: Bu komut SADECE ``VehicleFinancialProfile`` (finansman tipi + nakit/
kredi kırılımı) üretir. Kredi (``Loan``/``LoanVehicleAllocation``) oluşturma
buradan ÇIKARILDI — gerçek ``Kredi_Arac_Eslestirme.xlsx`` eşleştirmesini
kullanan ``reseed_vehicle_loans`` komutuna devredildi (bu komutun eski
"şirket başına sabit sayıda gruba böl" sezgisi, filoya erken giren küçük
paylı araçların olgunlaşmış bir kredinin tüm faizini üstlenmesi gibi ciddi
bir hataya yol açmıştı — artık kullanılmıyor). Bu iki komut BAĞIMSIZDIR;
``seed_real_vehicle_financials``'ı tekrar çalıştırmak kredilere dokunmaz.

Kullanım::

    python manage.py seed_real_vehicle_financials ../datas/Arac_Finansman_Tipi.xlsx
"""

from __future__ import annotations

from decimal import Decimal
from typing import Any

from django.core.management.base import BaseCommand, CommandError, CommandParser
from django.db import transaction

from backend.apps.vehicle_finance.choices import FinancingType
from backend.apps.vehicle_finance.models import VehicleFinancialProfile
from backend.apps.vehicles.models import Vehicle

COL_PLATE = "Plaka"
COL_TYPE = "Finansman Tipi"
COL_LOAN_AMOUNT = "Kredi ile Finanse Edilen Tutar"
COL_CASH_AMOUNT = "Nakit (Özkaynak) ile Finanse Edilen Tutar"

TYPE_MAP = {
    "Kredi": FinancingType.LOAN,
    "Karma": FinancingType.MIXED,
    "Nakit": FinancingType.CASH,
}


class Command(BaseCommand):
    help = "Gerçek araçlara VehicleFinancialProfile (finansman tipi + tutar kırılımı) üretir."

    def add_arguments(self, parser: CommandParser) -> None:
        parser.add_argument("xlsx_path", type=str, help="Arac_Finansman_Tipi.xlsx yolu.")
        parser.add_argument(
            "--dry-run", action="store_true", help="Hiçbir şey kaydetmeden özet göster."
        )

    def handle(self, *args: Any, **options: Any) -> None:
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError("openpyxl kurulu değil.") from exc

        xlsx_path: str = options["xlsx_path"]
        try:
            workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        except (OSError, KeyError) as exc:
            raise CommandError(f"Dosya okunamadı: {xlsx_path} ({exc})") from exc

        sheet = workbook[workbook.sheetnames[0]]
        rows_iter = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
        for required in (COL_PLATE, COL_TYPE, COL_LOAN_AMOUNT, COL_CASH_AMOUNT):
            if required not in header:
                raise CommandError(f"Eksik sütun: {required}")

        by_plate: dict[str, dict[str, object]] = {}
        for row in rows_iter:
            record = dict(zip(header, row))
            plate = str(record.get(COL_PLATE) or "").strip()
            if plate:
                by_plate[plate] = record

        vehicles = list(Vehicle.objects.select_related("financial_profile", "company").all())
        self.stdout.write(f"Toplam araç: {len(vehicles)}, kaynakta plaka: {len(by_plate)}")

        profiles_to_create: list[VehicleFinancialProfile] = []
        skipped_no_price = 0

        for vehicle in vehicles:
            if hasattr(vehicle, "financial_profile"):
                continue  # zaten profili var, dokunma
            if vehicle.purchase_price is None:
                skipped_no_price += 1
                continue

            source = by_plate.get(vehicle.plate_number)
            if source is not None:
                financing_type = TYPE_MAP.get(str(source.get(COL_TYPE) or "").strip(), FinancingType.CASH)
                cash_amount = Decimal(str(source.get(COL_CASH_AMOUNT) or 0))
            else:
                financing_type = FinancingType.CASH
                cash_amount = vehicle.purchase_price

            purchase_date = vehicle.purchase_invoice_date or vehicle.fleet_entry_date
            if purchase_date is None:
                skipped_no_price += 1
                continue

            profiles_to_create.append(
                VehicleFinancialProfile(
                    vehicle=vehicle,
                    purchase_date=purchase_date,
                    purchase_price=vehicle.purchase_price,
                    financing_type=financing_type,
                    cash_financed_amount=cash_amount,
                )
            )

        self.stdout.write(f"Oluşturulacak finansal profil: {len(profiles_to_create)}")
        if skipped_no_price:
            self.stdout.write(
                self.style.WARNING(f"{skipped_no_price} araç fiyat/tarih eksikliği yüzünden atlandı.")
            )

        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("--dry-run: hiçbir şey kaydedilmedi."))
            return

        with transaction.atomic():
            VehicleFinancialProfile.objects.bulk_create(profiles_to_create)

        self.stdout.write(self.style.SUCCESS(f"{len(profiles_to_create)} finansal profil oluşturuldu."))
        self.stdout.write(
            self.style.WARNING(
                "Kredi tahsisi için ayrıca: python manage.py reseed_vehicle_loans ..."
            )
        )
