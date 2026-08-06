"""TSB Kasko Değer Listesi'ni (xlsx) CarValuation tablosuna aktarır.

Kullanım:
    python manage.py import_kasko_values                # backend/202607R1.xlsx
    python manage.py import_kasko_values --file <yol>   # farklı dönem dosyası

Dosya formatı (TSB aylık liste, tek sayfa):
    1. satır: dönem başlığı (ör. "Temmuz 2026")
    2. satır: Marka Kodu | Tip Kodu | Marka Adı | Tip Adı | 2026 | 2025 | ...
    3+ satır: veriler; yıl hücresi 0 ise o yıl için kasko bedeli yoktur.

Import idempotenttir: tablo tamamen boşaltılıp dosyadan yeniden doldurulur.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from backend.apps.enuygun.models import CarValuation

DEFAULT_FILE = Path(settings.BASE_DIR) / "202607R1.xlsx"
BATCH_SIZE = 5_000


class Command(BaseCommand):
    help = "TSB Kasko Değer Listesi'ni (xlsx) CarValuation tablosuna aktarır."

    def add_arguments(self, parser) -> None:
        parser.add_argument(
            "--file",
            default=str(DEFAULT_FILE),
            help=f"TSB xlsx dosya yolu (varsayılan: {DEFAULT_FILE})",
        )

    def handle(self, *args, **options) -> None:
        path = Path(options["file"])
        if not path.exists():
            raise CommandError(f"Dosya bulunamadı: {path}")

        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise CommandError("openpyxl kurulu değil (bkz. requirements.txt).") from exc

        workbook = load_workbook(path, read_only=True)
        sheet = workbook.active

        rows = sheet.iter_rows(min_row=2, values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise CommandError("Dosya boş: başlık satırı bulunamadı.") from exc

        year_columns = self._parse_year_columns(header)
        if not year_columns:
            raise CommandError(
                "Başlıkta yıl sütunu bulunamadı; TSB liste formatı bekleniyor "
                "(Marka Kodu | Tip Kodu | Marka Adı | Tip Adı | 2026 | ...)."
            )

        created = self._import_rows(rows, year_columns)
        workbook.close()

        self.stdout.write(
            self.style.SUCCESS(
                f"{path.name}: {created:,} kasko değeri yüklendi "
                f"({len(year_columns)} yıl sütunu: "
                f"{max(year_columns.values())}–{min(year_columns.values())})."
            )
        )

    @staticmethod
    def _parse_year_columns(header: tuple) -> dict[int, int]:
        """Başlıktaki yıl sütunlarını {sütun_indeksi: yıl} olarak döndürür."""
        year_columns: dict[int, int] = {}
        for index, value in enumerate(header):
            if isinstance(value, int) and 1990 <= value <= 2100:
                year_columns[index] = value
        return year_columns

    @transaction.atomic
    def _import_rows(self, rows, year_columns: dict[int, int]) -> int:
        CarValuation.objects.all().delete()

        batch: list[CarValuation] = []
        created = 0

        for row in rows:
            brand_code, type_code = row[0], row[1]
            brand_name = (row[2] or "").strip() if isinstance(row[2], str) else ""
            type_name = (row[3] or "").strip() if isinstance(row[3], str) else ""
            if not brand_name or not type_name:
                continue
            if not isinstance(brand_code, int) or not isinstance(type_code, int):
                continue

            for index, year in year_columns.items():
                raw_value = row[index] if index < len(row) else None
                if not raw_value:  # None veya 0 → o yıl için değer yok
                    continue
                batch.append(
                    CarValuation(
                        brand_code=brand_code,
                        type_code=type_code,
                        brand_name=brand_name,
                        type_name=type_name,
                        year=year,
                        kasko_value=Decimal(str(raw_value)),
                    )
                )
                if len(batch) >= BATCH_SIZE:
                    CarValuation.objects.bulk_create(batch)
                    created += len(batch)
                    batch = []

        if batch:
            CarValuation.objects.bulk_create(batch)
            created += len(batch)

        return created
