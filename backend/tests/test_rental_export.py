"""GET /api/rentals/export/ — toplu kiralama Excel raporu view testi."""

from __future__ import annotations

from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from backend.apps.companies.models import Company, CustomUser, UserRole
from backend.apps.customer_profiles.models import Driver
from backend.apps.rentals.models import Rental, RentalStatus, ReturnStatus
from backend.apps.vehicles.models import Vehicle, VehicleStatus

pytestmark = pytest.mark.django_db

EXPORT_URL = "/api/rentals/export/"

EXPECTED_HEADERS = [
    "Sürücü",
    "Plaka",
    "Marka / Model",
    "Alış Tarihi",
    "Bırakış Tarihi",
    "Dönem",
    "Günlük Fiyat (₺)",
    "Toplam Fiyat (₺)",
    "Fiyat Farkı (₺)",
    "Teslim Durumu",
]


@pytest.fixture
def client(company: Company) -> APIClient:
    user = CustomUser.objects.create_user(
        username="test_personel_rez_export",
        password="test-pass-123",
        company=company,
        role=UserRole.OPERATIONS_STAFF,
    )
    api_client = APIClient()
    api_client.force_authenticate(user=user)
    return api_client


@pytest.fixture
def driver(company: Company) -> Driver:
    return Driver.objects.create(
        company=company,
        first_name="İhsan",
        last_name="Işık",
        national_id="66666666660",
        birth_date=timezone.now().date() - timedelta(days=365 * 40),
        phone="5551112233",
        license_number="R000222",
        city="İzmir",
    )


@pytest.fixture
def vehicle(company: Company) -> Vehicle:
    return Vehicle.objects.create(
        company=company,
        plate_number="35REZ100",
        brand="Fiat",
        model="Egea",
        production_year=2024,
        current_km=12_000,
        status=VehicleStatus.RENTED,
    )


def _make_rental(driver: Driver, vehicle: Vehicle, **overrides) -> Rental:
    now = timezone.now()
    defaults = dict(
        driver=driver,
        vehicle=vehicle,
        status=RentalStatus.ACTIVE,
        started_at=now - timedelta(days=2),
        ended_at=now + timedelta(days=8),
        start_km=vehicle.current_km,
        daily_rate=Decimal("1500.00"),
        quoted_total=Decimal("15000.00"),
    )
    defaults.update(overrides)
    return Rental.objects.create(**defaults)


def _load_sheet(response):
    return load_workbook(BytesIO(response.content)).active


def test_export_headers_and_row_content(client, driver, vehicle) -> None:
    _make_rental(driver, vehicle)

    response = client.get(EXPORT_URL)

    assert response.status_code == status.HTTP_200_OK
    assert "Kiralama_Raporu_" in response["Content-Disposition"]

    sheet = _load_sheet(response)
    assert [cell.value for cell in sheet[1]] == EXPECTED_HEADERS
    assert sheet.max_row == 2

    row = [cell.value for cell in sheet[2]]
    assert row[0] == "İhsan Işık"
    assert row[1] == "35REZ100"
    assert row[2] == "Fiat Egea"
    assert row[5] == "Aktif"
    assert float(row[6]) == 1500.0
    assert float(row[7]) == 15000.0
    assert row[9] == "—"  # henüz teslim alınmadı


def test_export_period_and_delivery_filters(client, driver, vehicle) -> None:
    now = timezone.now()
    _make_rental(driver, vehicle)  # AKTIF
    _make_rental(
        driver,
        vehicle,
        started_at=now + timedelta(days=3),
        ended_at=now + timedelta(days=6),
    )  # GELECEK
    _make_rental(
        driver,
        vehicle,
        status=RentalStatus.COMPLETED,
        started_at=now - timedelta(days=20),
        ended_at=now - timedelta(days=10),
        actual_return_at=now - timedelta(days=12),
        return_status=ReturnStatus.EARLY,
        price_difference=Decimal("-3000.00"),
    )  # PAST / erken teslim

    response = client.get(EXPORT_URL, {"period": "UPCOMING"})
    sheet = _load_sheet(response)
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=6).value == "Gelecek"

    response = client.get(EXPORT_URL, {"period": "PAST", "return_timing": "EARLY"})
    sheet = _load_sheet(response)
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=10).value == "Erken Teslim"
    assert float(sheet.cell(row=2, column=9).value) == -3000.0

    response = client.get(EXPORT_URL, {"period": "PAST", "return_timing": "LATE"})
    sheet = _load_sheet(response)
    assert sheet.max_row == 1  # yalnızca başlık


def test_export_does_not_leak_other_companys_records(client, driver, vehicle) -> None:
    _make_rental(driver, vehicle)

    rival = Company.objects.create(name="rez-rival", tax_number="8880005555")
    rival_driver = Driver.objects.create(
        company=rival,
        first_name="Rakip",
        last_name="Kişi",
        national_id="77777777770",
        birth_date=timezone.now().date() - timedelta(days=365 * 30),
        phone="5554443322",
        license_number="R000333",
        city="Bursa",
    )
    rival_vehicle = Vehicle.objects.create(
        company=rival,
        plate_number="34RKP200",
        brand="Renault",
        model="Clio",
        production_year=2023,
        current_km=5_000,
        status=VehicleStatus.RENTED,
    )
    _make_rental(rival_driver, rival_vehicle)

    sheet = _load_sheet(client.get(EXPORT_URL))
    plates = [sheet.cell(row=r, column=2).value for r in range(2, sheet.max_row + 1)]
    assert plates == ["35REZ100"]
