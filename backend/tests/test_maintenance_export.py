"""GET /api/vehicles/maintenance-alerts/export/ — bakım Excel raporu testi.

Bakım durumu artık aracın SON BAKIM kaydından (VehicleExpense, kategori
VEHICLE_MAINTENANCE) itibaren km/gün'e göre hesaplanıyor — bare
``current_km`` tek başına artık durumu belirlemiyor, bu yüzden her araç
gerçek bir son-bakım referans noktasıyla kurulur (bkz. maintenance_service.py).
"""

from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from backend.apps.companies.models import Company, CustomUser, UserRole
from backend.apps.vehicle_finance.choices import ExpenseCategory
from backend.apps.vehicle_finance.models import VehicleExpense
from backend.apps.vehicles.models import Vehicle, VehicleStatus

pytestmark = pytest.mark.django_db

EXPORT_URL = "/api/vehicles/maintenance-alerts/export/"

EXPECTED_HEADERS = [
    "Plaka",
    "Marka / Model",
    "Bakım Durumu",
    "Güncel KM",
    "Son Bakım Tarihi",
    "Son Bakım KM",
    "Hedef KM",
    "Kalan KM",
    "İlerleme (%)",
    "Hedef Tarih",
    "Kalan Gün",
]

_RECENT = date.today() - timedelta(days=5)


@pytest.fixture
def client(company: Company) -> APIClient:
    user = CustomUser.objects.create_user(
        username="test_personel_bakim_export",
        password="test-pass-123",
        company=company,
        role=UserRole.OPERATIONS_STAFF,
    )
    api_client = APIClient()
    api_client.force_authenticate(user=user)
    return api_client


def _make_vehicle(
    company: Company,
    plate: str,
    current_km: int,
    *,
    last_maintenance_km: int = 0,
    last_maintenance_date: date = _RECENT,
) -> Vehicle:
    """Bir araç + gerçek bir 'son bakım' gideri oluşturur (km/tarih referansı)."""
    vehicle = Vehicle.objects.create(
        company=company,
        plate_number=plate,
        brand="Toyota",
        model="Corolla",
        production_year=2023,
        current_km=current_km,
        status=VehicleStatus.AVAILABLE,
    )
    VehicleExpense.objects.create(
        vehicle=vehicle,
        category=ExpenseCategory.VEHICLE_MAINTENANCE,
        amount="500.00",
        expense_date=last_maintenance_date,
        odometer_km=last_maintenance_km,
    )
    return vehicle


def _load_sheet(response):
    return load_workbook(BytesIO(response.content)).active


def test_export_pending_maintenance(client, company) -> None:
    # Son bakımdan bu yana 14.500 km → 10.000-15.000 penceresinde → WARNING.
    _make_vehicle(company, "06BKM001", 14_500)
    # Son bakımdan bu yana 2.000 km → sağlıklı, uyarı listesine girmez.
    _make_vehicle(company, "06SGL001", 2_000)

    response = client.get(EXPORT_URL)

    assert response.status_code == status.HTTP_200_OK
    assert "Bakim_Raporu_" in response["Content-Disposition"]

    sheet = _load_sheet(response)
    assert [cell.value for cell in sheet[1]] == EXPECTED_HEADERS
    assert sheet.max_row == 2

    row = [cell.value for cell in sheet[2]]
    assert row[0] == "06BKM001"
    assert row[2] == "Yaklaşıyor"
    assert row[3] == 14_500
    assert row[5] == 0  # Son Bakım KM
    assert row[6] == 15_000  # Hedef KM
    assert row[7] == 500  # Kalan KM


def test_export_pending_excludes_critical_vehicles(client, company) -> None:
    """Regresyon: 'Bakım Bekleyen' (PENDING) yalnızca WARNING içermeli,
    'Acil' (CRITICAL) ile aynı araçları LİSTELEMEMELİ."""
    _make_vehicle(company, "06BKM005", 14_500)  # WARNING
    _make_vehicle(company, "06ACL002", 15_500)  # CRITICAL

    response = client.get(EXPORT_URL)  # PENDING (varsayılan)
    sheet = _load_sheet(response)

    plates = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
    assert plates == ["06BKM005"]


def test_export_healthy_filter(client, company) -> None:
    _make_vehicle(company, "06BKM002", 14_500)
    _make_vehicle(company, "06SGL002", 2_000)

    response = client.get(EXPORT_URL, {"status_filter": "HEALTHY"})
    sheet = _load_sheet(response)

    assert sheet.max_row == 2
    row = [cell.value for cell in sheet[2]]
    assert row[0] == "06SGL002"
    assert row[2] == "Sağlıklı"
    assert row[3] == 2_000
    assert row[6] == 15_000  # sağlıklı araçta da km referansı biliniyorsa hedef km hesaplanır


def test_export_urgent_filter(client, company) -> None:
    _make_vehicle(company, "06BKM003", 14_500)  # WARNING (10k-15k arası)
    _make_vehicle(company, "06ACL001", 15_200)  # CRITICAL (>= 15.000)

    response = client.get(EXPORT_URL, {"status_filter": "URGENT"})
    sheet = _load_sheet(response)
    assert sheet.max_row == 2  # yalnızca acil araç
    row = [cell.value for cell in sheet[2]]
    assert row[0] == "06ACL001"
    assert row[2] == "Acil"


def test_export_does_not_leak_other_companys_vehicles(client, company) -> None:
    _make_vehicle(company, "06BKM004", 14_500)
    rival = Company.objects.create(name="bakim-rival", tax_number="8880007777")
    _make_vehicle(rival, "34RKP300", 14_500)

    sheet = _load_sheet(client.get(EXPORT_URL))
    plates = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
    assert plates == ["06BKM004"]
