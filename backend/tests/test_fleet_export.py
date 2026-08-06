"""GET /api/vehicles/export/ — toplu filo Excel raporu view testi."""

from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from backend.apps.companies.models import Company, CustomUser, UserRole
from backend.apps.vehicle_finance.choices import FinancingType, InsurancePolicyType
from backend.apps.vehicle_finance.models import (
    CompanyOverheadExpense,
    InsurancePolicy,
    RentalIncome,
    VehicleFinancialProfile,
)
from backend.apps.vehicles.models import Vehicle, VehicleStatus

pytestmark = pytest.mark.django_db

EXPORT_URL = "/api/vehicles/export/"

EXPECTED_HEADERS = [
    "Plaka",
    "Marka / Model",
    "Durum",
    "Model Yılı",
    "Kilometre",
    "Yakıt Tipi",
    "Şanzıman",
    "Sınıf / SIPP Kodu",
    "Alış Fiyatı (₺)",
    "Günlük Maliyet (₺)",
    "Biriken Net Kâr (₺)",
]


@pytest.fixture
def other_company() -> Company:
    return Company.objects.create(name="rival-tenant", tax_number="8880003333")


@pytest.fixture
def staff(company: Company) -> CustomUser:
    return CustomUser.objects.create_user(
        username="test_personel_export",
        password="test-pass-123",
        company=company,
        role=UserRole.OPERATIONS_STAFF,
    )


@pytest.fixture
def client(staff: CustomUser) -> APIClient:
    api_client = APIClient()
    api_client.force_authenticate(user=staff)
    return api_client


def _make_vehicle(company: Company, plate: str, **overrides) -> Vehicle:
    defaults = dict(
        company=company,
        plate_number=plate,
        brand="Renault",
        model="Clio",
        production_year=2023,
        current_km=42_000,
        status=VehicleStatus.AVAILABLE,
        fuel_type="GASOLINE",
        transmission="MANUAL",
        purchase_price=Decimal("900000.00"),
        initial_industrial_expense=Decimal("10000.00"),
        annual_kasko_cost=Decimal("30000.00"),
        annual_traffic_insurance_cost=Decimal("5000.00"),
        annual_mtv_cost=Decimal("9000.00"),
        estimated_monthly_maintenance=Decimal("2500.00"),
        fleet_entry_date=date.today() - timedelta(days=200),
    )
    defaults.update(overrides)
    return Vehicle.objects.create(**defaults)


def _load_sheet(response):
    return load_workbook(BytesIO(response.content)).active


def test_export_includes_all_vehicles_and_headers(client: APIClient, company: Company) -> None:
    vehicle = _make_vehicle(company, "06EXP001")
    _make_vehicle(company, "06EXP002", status=VehicleStatus.RENTED)
    # `_make_vehicle` her zaman bu alanları doldurur (bkz. defaults) — model
    # alanları nullable olduğu için mypy bunu bilemez.
    assert vehicle.fleet_entry_date is not None
    assert vehicle.purchase_price is not None

    # Günlük maliyet/biriken kâr sütunlarının hesaplanabilmesi için gerçek
    # maliyet motorunun (§3) ihtiyaç duyduğu asgari veri: finansman profili,
    # sigorta poliçesi, genel gider kaydı.
    VehicleFinancialProfile.objects.create(
        vehicle=vehicle,
        purchase_date=vehicle.fleet_entry_date,
        purchase_price=vehicle.purchase_price,
        financing_type=FinancingType.CASH,
        cash_financed_amount=vehicle.purchase_price,
    )
    InsurancePolicy.objects.create(
        vehicle=vehicle,
        policy_type=InsurancePolicyType.COMPREHENSIVE,
        amount=Decimal("20000.00"),
        start_date=vehicle.fleet_entry_date,
        end_date=vehicle.fleet_entry_date + timedelta(days=365),
    )
    CompanyOverheadExpense.objects.create(
        company=company, expense_date=date.today(), amount=Decimal("1000.00")
    )

    RentalIncome.objects.create(
        vehicle=vehicle,
        start_date=date(2026, 6, 1),
        end_date=date(2026, 6, 11),  # 10 gün
        base_amount=Decimal("50000.00"),
    )

    response = client.get(EXPORT_URL)

    assert response.status_code == status.HTTP_200_OK
    assert (
        response["Content-Type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "Filo_Envanter_Raporu_" in response["Content-Disposition"]
    assert ".xlsx" in response["Content-Disposition"]

    sheet = _load_sheet(response)
    headers = [cell.value for cell in sheet[1]]
    assert headers == EXPECTED_HEADERS
    assert sheet.max_row == 3  # başlık + 2 araç

    first_row = [cell.value for cell in sheet[2]]
    assert first_row[0] == "06EXP001"
    assert first_row[1] == "Renault Clio"
    assert first_row[2] == "Müsait"
    assert first_row[5] == "Benzin"
    assert first_row[6] == "Manuel"
    assert first_row[7].endswith("ECMR")
    # Günlük maliyet: (900k + 10k + yıllık işletme) / 365 — pozitif olmalı.
    assert first_row[9] > 0
    # Biriken kâr = 50.000 gelir − 10 gün × günlük maliyet.
    assert first_row[10] == pytest.approx(50000 - 10 * first_row[9], abs=0.05)


def test_export_filters_by_status_and_search(client: APIClient, company: Company) -> None:
    _make_vehicle(company, "06FIL001", status=VehicleStatus.RENTED, brand="Toyota", model="Corolla")
    _make_vehicle(company, "06FIL002", status=VehicleStatus.AVAILABLE)

    response = client.get(EXPORT_URL, {"status": "RENTED"})
    sheet = _load_sheet(response)
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=1).value == "06FIL001"

    # Türkçe uyumlu, harf duyarsız arama (İ → i).
    response = client.get(EXPORT_URL, {"q": "COROLLA"})
    sheet = _load_sheet(response)
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=2).value == "Toyota Corolla"

    response = client.get(EXPORT_URL, {"q": "yok-boyle-arac"})
    sheet = _load_sheet(response)
    assert sheet.max_row == 1  # yalnızca başlık


def test_export_does_not_leak_other_companys_vehicles(
    client: APIClient, company: Company, other_company: Company
) -> None:
    _make_vehicle(company, "06BIZ001")
    _make_vehicle(other_company, "34RKP001")

    response = client.get(EXPORT_URL)
    sheet = _load_sheet(response)
    plates = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
    assert plates == ["06BIZ001"]
